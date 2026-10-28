from __future__ import annotations

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID
from zipfile import ZIP_STORED, ZipFile

import pytest
import httpx
from typing import Any
from docx import Document
from fastapi import HTTPException
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.extraction import (
    DocumentProcessingError,
    chunk_text,
    extract_docx_chunks,
    extract_doc_chunks,
    extract_epub_chunks,
    extract_image_chunks,
    extract_pdf_chunks,
    extract_pptx_chunks,
    extract_ppt_chunks,
)
from app.auth import AuthenticatedUser
from app.config import Settings
from app.upload import courses as courses_module
from app.upload.courses import CourseExtractionService, get_course_extraction_service
from app.extraction import image as image_module
from app.extraction import legacy_office as legacy_office_module
from app.extraction import pdf as pdf_module


def _pdf_with_text(*page_texts: str) -> bytes:
    page_numbers = range(3, 3 + len(page_texts))
    font_number = 3 + len(page_texts)
    first_content_number = font_number + 1
    page_references = " ".join(f"{number} 0 R" for number in page_numbers)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{page_references}] /Count {len(page_texts)} >>".encode("ascii"),
    ]
    for index in range(len(page_texts)):
        page_number = 3 + index
        content_number = first_content_number + index
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_number} 0 R >> >> /Contents {content_number} 0 R >>".encode("ascii")
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in page_texts:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        objects.append(f"<< /Length {len(content)} >>\nstream\n".encode("ascii") + content + b"\nendstream")

    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_number} 0 obj\n".encode("ascii"))
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets:
        pdf.extend(f"{offset:010} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    )
    return bytes(pdf)


def _docx_with_content() -> bytes:
        document = Document()
        document.add_heading("Contexte", level=1)
        document.add_paragraph("La lumière naturelle organise les espaces.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Orientation"
        table.cell(0, 1).text = "Sud"
        output = BytesIO()
        document.save(output)
        return output.getvalue()


def _pptx_with_content() -> bytes:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        textbox.text = "Diapositive sur la ventilation"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "Débit"
        table.cell(0, 1).text = "Contrôlé"
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()


def _epub_with_chapters() -> bytes:
        output = BytesIO()
        with ZipFile(output, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip", compress_type=ZIP_STORED)
                archive.writestr("META-INF/container.xml", """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
    <rootfiles><rootfile full-path="OEBPS/book.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>""")
                archive.writestr("OEBPS/book.opf", """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
    <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">test-book</dc:identifier><dc:title>Test</dc:title><dc:language>fr</dc:language></metadata>
    <manifest><item id="one" href="one.xhtml" media-type="application/xhtml+xml"/><item id="two" href="two.xhtml" media-type="application/xhtml+xml"/></manifest>
    <spine><itemref idref="one"/><itemref idref="two"/></spine>
</package>""")
                archive.writestr("OEBPS/one.xhtml", "<html><body><h1>Chapitre un</h1><p>Premier chapitre sur le site.</p></body></html>")
                archive.writestr("OEBPS/two.xhtml", "<html><body><h1>Chapitre deux</h1><p>Deuxième chapitre sur la structure.</p></body></html>")
        return output.getvalue()


def _png_image() -> bytes:
    output = BytesIO()
    Image.new("RGB", (40, 30), "white").save(output, format="PNG")
    return output.getvalue()


def _legacy_office_file() -> bytes:
    return bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy test payload"


def test_extract_pdf_text_preserves_page_and_chunk_references() -> None:
    chunks = extract_pdf_chunks(
        _pdf_with_text("Architecture course reference supports spatial design choices"),
        max_chars=24,
        overlap_chars=6,
    )

    assert len(chunks) > 1
    assert all(chunk.page_number == 1 for chunk in chunks)
    assert all(chunk.locator_type == "page" and chunk.locator == "Page 1" for chunk in chunks)
    assert all(chunk.source_order == 1 for chunk in chunks)
    assert [chunk.chunk_index for chunk in chunks] == list(range(1, len(chunks) + 1))
    assert all(chunk.text for chunk in chunks)
    assert all(len(chunk.text) <= 24 for chunk in chunks)


def test_extract_pdf_chunks_keep_their_source_page() -> None:
    chunks = extract_pdf_chunks(
        _pdf_with_text("first page content", "second page content"),
        max_chars=64,
        overlap_chars=0,
    )

    assert [(chunk.page_number, chunk.chunk_index) for chunk in chunks] == [(1, 1), (2, 1)]
    assert [(chunk.locator, chunk.source_order) for chunk in chunks] == [("Page 1", 1), ("Page 2", 2)]
    assert [chunk.text for chunk in chunks] == ["first page content", "second page content"]


def test_extract_pdf_runs_ocr_only_for_pages_without_native_text(monkeypatch: pytest.MonkeyPatch) -> None:
    recognized_pages: list[tuple[int, int, int]] = []

    def recognize(image: np.ndarray) -> str:
        recognized_pages.append(image.shape)
        return "Texte lu sur le scan"

    monkeypatch.setattr(pdf_module, "recognize_text", recognize)
    chunks = extract_pdf_chunks(_pdf_with_text("", "Cette page contient assez de texte natif pour eviter OCR."))

    assert len(chunks) == 2
    assert chunks[0].text == "Texte lu sur le scan"
    assert chunks[0].page_number == 1
    assert chunks[1].page_number == 2
    assert len(recognized_pages) == 1


def test_extract_image_chunks_uses_ocr_and_image_provenance(monkeypatch: pytest.MonkeyPatch) -> None:
    image_buffer = BytesIO()
    Image.new("RGB", (40, 30), "white").save(image_buffer, format="PNG")
    received_shapes: list[tuple[int, ...]] = []

    def recognize(image: np.ndarray) -> str:
        received_shapes.append(image.shape)
        return "Plan de ventilation"

    monkeypatch.setattr(image_module, "recognize_text", recognize)
    chunks = extract_image_chunks(image_buffer.getvalue(), "png")

    assert len(chunks) == 1
    assert chunks[0].text == "Plan de ventilation"
    assert chunks[0].page_number is None
    assert chunks[0].locator_type == "image"
    assert chunks[0].locator == "Image"
    assert received_shapes == [(30, 40, 3)]


def test_extract_docx_keeps_sections_and_tables_without_fake_pages() -> None:
    chunks = extract_docx_chunks(_docx_with_content())

    assert len(chunks) == 1
    assert chunks[0].page_number is None
    assert chunks[0].locator_type == "section"
    assert chunks[0].locator == "Section 1: Contexte"
    assert "La lumière naturelle" in chunks[0].text
    assert "Orientation | Sud" in chunks[0].text


def test_extract_pptx_keeps_slide_and_table_provenance() -> None:
    chunks = extract_pptx_chunks(_pptx_with_content())

    assert len(chunks) == 1
    assert chunks[0].page_number is None
    assert chunks[0].locator_type == "slide"
    assert chunks[0].locator == "Diapositive 1"
    assert "ventilation" in chunks[0].text
    assert "Débit | Contrôlé" in chunks[0].text


def test_extract_epub_follows_spine_order_and_chapter_names() -> None:
    chunks = extract_epub_chunks(_epub_with_chapters())

    assert [(chunk.source_order, chunk.locator_type) for chunk in chunks] == [(1, "chapter"), (2, "chapter")]
    assert [chunk.locator for chunk in chunks] == [
        "Chapitre 1: Chapitre un (one.xhtml)",
        "Chapitre 2: Chapitre deux (two.xhtml)",
    ]
    assert [chunk.text for chunk in chunks] == [
        "Chapitre un Premier chapitre sur le site.",
        "Chapitre deux Deuxième chapitre sur la structure.",
    ]


@pytest.mark.parametrize("extractor", [extract_docx_chunks, extract_pptx_chunks, extract_epub_chunks])
def test_office_and_epub_parsers_reject_invalid_archives(extractor) -> None:
    with pytest.raises(DocumentProcessingError, match="invalid|incomplete"):
        extractor(b"not a ZIP package")


def test_chunk_text_returns_empty_for_whitespace() -> None:
    assert chunk_text("  \n\t  ") == []


@pytest.mark.parametrize(
    ("max_chars", "overlap_chars"),
    [(0, 0), (10, -1), (10, 10)],
)
def test_chunk_text_rejects_invalid_limits(max_chars: int, overlap_chars: int) -> None:
    with pytest.raises(ValueError):
        chunk_text("some text", max_chars=max_chars, overlap_chars=overlap_chars)


@pytest.mark.parametrize(
    ("pdf_bytes", "message"),
    [
        (b"", "empty"),
        (b"not a pdf", "invalid"),
        (_pdf_with_text(""), "no extractable text"),
    ],
)
def test_extract_pdf_rejects_empty_invalid_or_scanned_documents(pdf_bytes: bytes, message: str) -> None:
    with pytest.raises(DocumentProcessingError, match=message):
        extract_pdf_chunks(pdf_bytes)


class FakeCourseClient:
    def __init__(self, content: bytes, content_type: str = "application/pdf") -> None:
        self.content = content
        self.rows = [{
            "id": "00000000-0000-0000-0000-000000000001",
            "storage_path": "user-a/course.pdf",
            "content_type": content_type,
            "file_size": len(content),
        }]
        self.filters = []
        self.saved: dict[str, Any] = {}
        self.downloads = []
        self.storage = self

    def table(self, table_name):
        assert table_name == "user_courses"
        return self

    def select(self, columns):
        return self

    def eq(self, column, value):
        self.filters.append((column, value))
        return self

    def execute(self):
        return SimpleNamespace(data=self.rows)

    def from_(self, bucket):
        assert bucket == "user-courses"
        return self

    def download(self, path):
        self.downloads.append(path)
        return self.content

    def rpc(self, fn, params=None):
        assert fn == "save_course_extraction"
        assert params is not None
        self.saved = params
        return self


def test_course_extraction_persists_pages_and_uses_owner_filter() -> None:
    client = FakeCourseClient(_pdf_with_text("first page", "second page"))
    course_id = UUID(client.rows[0]["id"])

    result = CourseExtractionService(client, "user-a").extract(course_id)

    assert result.status == "extracted"
    assert result.chunk_count == 2
    assert client.filters == [("id", str(course_id)), ("owner_id", "user-a")]
    assert client.saved["result_chunks"] == [
        {"text": "first page", "page_number": 1, "chunk_index": 1,
         "locator_type": "page", "locator": "Page 1", "source_order": 1, "document_format": "pdf"},
        {"text": "second page", "page_number": 2, "chunk_index": 1,
         "locator_type": "page", "locator": "Page 2", "source_order": 2, "document_format": "pdf"},
    ]


@pytest.mark.parametrize(
    ("content_type", "content", "document_format", "locator_type"),
    [
        ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", _docx_with_content(), "docx", "section"),
        ("application/vnd.openxmlformats-officedocument.presentationml.presentation", _pptx_with_content(), "pptx", "slide"),
        ("application/epub+zip", _epub_with_chapters(), "epub", "chapter"),
    ],
    ids=["docx", "pptx", "epub"],
)
def test_course_extraction_dispatches_modern_formats_and_persists_locators(
    content_type: str,
    content: bytes,
    document_format: str,
    locator_type: str,
) -> None:
    client = FakeCourseClient(content, content_type)

    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))

    assert result.status == "extracted"
    assert result.chunk_count > 0
    assert client.saved["result_chunks"]
    assert all(chunk["document_format"] == document_format for chunk in client.saved["result_chunks"])
    assert all(chunk["locator_type"] == locator_type for chunk in client.saved["result_chunks"])
    assert all(chunk["page_number"] is None for chunk in client.saved["result_chunks"])
    assert client.downloads == [client.rows[0]["storage_path"]]


@pytest.mark.parametrize("content", [b"invalid pdf", _pdf_with_text("")])
def test_course_extraction_records_parse_failure(content: bytes) -> None:
    client = FakeCourseClient(content)
    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))
    assert result.status == "failed"
    assert result.error
    assert client.saved["result_status"] == "failed"
    assert client.saved["result_chunks"] == []


def test_course_extraction_leaves_other_formats_stored_without_parsing() -> None:
    client = FakeCourseClient(b"unknown content", "application/octet-stream")
    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))
    assert result.status == "unsupported"
    assert client.downloads == []


def test_course_extraction_dispatches_image_ocr_and_persists_locator(monkeypatch: pytest.MonkeyPatch) -> None:
    client = FakeCourseClient(_png_image(), "image/png")
    monkeypatch.setattr(image_module, "recognize_text", lambda image: "Texte de plan reconnu")

    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))

    assert result.status == "extracted"
    assert client.saved["result_chunks"] == [{
        "text": "Texte de plan reconnu",
        "page_number": None,
        "chunk_index": 1,
        "locator_type": "image",
        "locator": "Image",
        "source_order": 1,
        "document_format": "png",
    }]


@pytest.mark.parametrize(
    ("extractor", "converted_content", "target_format", "locator_type"),
    [
        (extract_doc_chunks, _docx_with_content(), "docx", "section"),
        (extract_ppt_chunks, _pptx_with_content(), "pptx", "slide"),
    ],
    ids=["doc", "ppt"],
)
def test_legacy_office_conversion_uses_headless_isolated_profile(
    monkeypatch: pytest.MonkeyPatch,
    extractor,
    converted_content: bytes,
    target_format: str,
    locator_type: str,
) -> None:
    monkeypatch.setenv(legacy_office_module.LEGACY_CONVERSION_OPT_IN, "true")
    monkeypatch.setenv("LIBREOFFICE_BINARY", "fake-libreoffice")
    calls: list[dict[str, Any]] = []

    def run(command, **options):
        calls.append({"command": command, **options})
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / f"source.{target_format}").write_bytes(converted_content)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(legacy_office_module.subprocess, "run", run)
    chunks = extractor(_legacy_office_file())

    assert chunks[0].locator_type == locator_type
    assert "--headless" in calls[0]["command"]
    assert any(arg.startswith("-env:UserInstallation=file:") for arg in calls[0]["command"])
    assert calls[0]["timeout"] == legacy_office_module.CONVERSION_TIMEOUT_SECONDS
    assert calls[0]["capture_output"] is True


def test_legacy_office_conversion_fails_explicitly_without_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(legacy_office_module.LEGACY_CONVERSION_OPT_IN, "true")
    monkeypatch.delenv("LIBREOFFICE_BINARY", raising=False)
    monkeypatch.setattr(legacy_office_module.shutil, "which", lambda _: None)

    with pytest.raises(DocumentProcessingError, match="requires LibreOffice"):
        extract_doc_chunks(_legacy_office_file())


def test_legacy_office_conversion_is_disabled_without_explicit_opt_in(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv(legacy_office_module.LEGACY_CONVERSION_OPT_IN, raising=False)
    monkeypatch.setenv("LIBREOFFICE_BINARY", "fake-libreoffice")

    with pytest.raises(DocumentProcessingError, match="disabled until an isolated converter"):
        extract_doc_chunks(_legacy_office_file())


def test_course_extraction_reports_disabled_legacy_conversion(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv(legacy_office_module.LEGACY_CONVERSION_OPT_IN, raising=False)
    client = FakeCourseClient(_legacy_office_file(), "application/msword")

    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))

    assert result.status == "failed"
    assert result.chunk_count == 0
    assert "disabled until an isolated converter" in (result.error or "")


def test_legacy_office_mime_types_are_connected_to_converters() -> None:
    assert courses_module.EXTRACTORS["application/msword"] == ("doc", extract_doc_chunks)
    assert courses_module.EXTRACTORS["application/vnd.ms-powerpoint"] == ("ppt", extract_ppt_chunks)


def test_course_extraction_rejects_foreign_storage_path() -> None:
    client = FakeCourseClient(b"pdf")
    client.rows[0]["storage_path"] = "other-user/course.pdf"
    with pytest.raises(HTTPException) as raised:
        CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))
    assert raised.value.status_code == 403
    assert client.saved == {}
    assert client.downloads == []


def test_course_extraction_does_not_download_invisible_course() -> None:
    client = FakeCourseClient(b"pdf")
    client.rows = []
    with pytest.raises(HTTPException) as raised:
        CourseExtractionService(client, "user-b").extract(UUID(int=1))
    assert raised.value.status_code == 404
    assert client.saved == {}
    assert client.downloads == []


def test_course_extraction_rejects_oversized_file_before_download() -> None:
    client = FakeCourseClient(b"pdf")
    client.rows[0]["file_size"] = 50 * 1024 * 1024 + 1
    result = CourseExtractionService(client, "user-a").extract(UUID(client.rows[0]["id"]))
    assert result.status == "failed"
    assert client.downloads == []


def test_course_extraction_propagates_jwt_to_storage_and_database(monkeypatch: pytest.MonkeyPatch) -> None:
    course_id = UUID(int=1)
    pdf = _pdf_with_text("private architecture course")
    requests = []

    def send(client, request, **kwargs):
        requests.append(request)
        if request.url.path == "/rest/v1/user_courses":
            return httpx.Response(200, request=request, json=[{
                "id": str(course_id), "storage_path": "user-a/course.pdf",
                "content_type": "application/pdf", "file_size": len(pdf),
            }])
        if request.url.path == "/storage/v1/object/user-courses/user-a/course.pdf":
            return httpx.Response(200, request=request, content=pdf)
        assert request.url.path == "/rest/v1/rpc/save_course_extraction"
        return httpx.Response(200, request=request, json=[])

    monkeypatch.setattr(httpx.Client, "send", send)
    monkeypatch.setattr(courses_module, "get_settings", lambda: Settings(
        "https://example.supabase.co", "test.api.key", "chunks", "test-model",
    ))
    service = get_course_extraction_service(AuthenticatedUser(id="user-a", access_token="user-jwt"))

    result = service.extract(course_id)

    assert result.status == "extracted"
    assert len(requests) == 3
    assert all(request.headers["authorization"] == "Bearer user-jwt" for request in requests)
    assert requests[0].url.params["owner_id"] == "eq.user-a"