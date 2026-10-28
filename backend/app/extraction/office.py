from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import DocumentChunk, DocumentProcessingError, chunk_text
from .package import validate_office_archive


def extract_docx_chunks(data: bytes) -> list[DocumentChunk]:
    if not data:
        raise DocumentProcessingError("The DOCX is empty.")
    validate_office_archive(data, "word/document.xml")
    try:
        document = Document(BytesIO(data))
        sections: list[tuple[str, str]] = []
        section_title = "Document"
        section_parts: list[str] = []

        def finish_section() -> None:
            if section_parts:
                sections.append((section_title, "\n".join(section_parts)))
                section_parts.clear()

        for element in document.element.body.iterchildren():
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                if paragraph.style and (paragraph.style.name or "").startswith("Heading"):
                    finish_section()
                    section_title = text
                else:
                    section_parts.append(text)
            elif element.tag == qn("w:tbl"):
                table = Table(element, document)
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                section_parts.extend(row for row in rows if row.strip(" |"))
        finish_section()
    except Exception as exc:
        raise DocumentProcessingError("The DOCX is invalid or could not be read.") from exc
    return _chunks_for_sections(sections)


def extract_pptx_chunks(data: bytes) -> list[DocumentChunk]:
    if not data:
        raise DocumentProcessingError("The PPTX is empty.")
    validate_office_archive(data, "ppt/presentation.xml")
    try:
        presentation = Presentation(BytesIO(data))
        chunks: list[DocumentChunk] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = _shape_text(slide.shapes)
            chunks.extend(
                DocumentChunk(
                    page_number=None,
                    chunk_index=index,
                    text=text,
                    locator_type="slide",
                    locator=f"Diapositive {slide_number}",
                    source_order=slide_number,
                )
                for index, text in enumerate(chunk_text("\n".join(parts)), start=1)
            )
    except Exception as exc:
        raise DocumentProcessingError("The PPTX is invalid or could not be read.") from exc
    if not chunks:
        raise DocumentProcessingError("The PPTX contains no extractable text.")
    return chunks


def _chunks_for_sections(sections: list[tuple[str, str]]) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    for source_order, (title, text) in enumerate(sections, start=1):
        locator = f"Section {source_order}: {title}"[:500]
        chunks.extend(
            DocumentChunk(
                page_number=None,
                chunk_index=index,
                text=chunk,
                locator_type="section",
                locator=locator,
                source_order=source_order,
            )
            for index, chunk in enumerate(chunk_text(text), start=1)
        )
    if not chunks:
        raise DocumentProcessingError("The document contains no extractable text.")
    return chunks


def _shape_text(shapes) -> list[str]:
    parts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.extend(_shape_text(shape.shapes))
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        elif getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                parts.append(text)
    return parts